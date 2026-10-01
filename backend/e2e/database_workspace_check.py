"""Focused HTTP/SQL workflow. Only UUID-named disposable SQL objects are changed.
Run: .venv/Scripts/python.exe -m backend.e2e.database_workspace_check

Failure cases:
- Changing the viewer target never changes scheduler, labware, Restore or a pending
  operation. Two databases with different schemas (including duplicate table names)
  browse correctly; an unconfigured viewer or report never falls back to the writer.
- Remote or non-admin callers cannot change connections or select another database;
  removing the selected connection fails until it is reassigned.
- Account creation is reviewed first, never modifies an existing login, escapes
  identifiers, rolls back failed grants and never saves or logs the one-off
  administrator credentials. The new identity can SELECT but not write or run DDL.
- Editing an installed report keeps sibling tools, files, inputs and mappings; a
  changed installed base or mapping prevents publishing an old draft.
- Preview uses the same robot/scheduler gate as execution and rolls back. Report
  failure, timeout or exit releases its slot and leaves the backend responsive.
"""
from contextlib import nullcontext
import io
import json
from pathlib import Path
import tempfile
import time
import traceback
import zipfile
import os
import openpyxl
from fastapi import FastAPI, Request, HTTPException
from fastapi.testclient import TestClient
from backend.e2e.report_wizard_check import sql_fixture, ROOT, BASE
from backend.api.database_tools import router
from backend.api.database import router as viewer_router
from backend.services.auth import get_current_user
from backend.services.database_tools import DatabaseTools, get_database_tools

EVIDENCE = Path(os.environ.get('ROBOTCONTROL_E2E_EVIDENCE', str(ROOT / 'test-output/database-workspace-verification')))


def run():
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    result = dict(passed=False, command='.venv/Scripts/python.exe -m backend.e2e.database_workspace_check', checks=[])
    try:
        with sql_fixture() as fixture, tempfile.TemporaryDirectory(prefix='rc-workspace-') as temp:
            admin = fixture['admin']; new_login = fixture['login'] + '_reader'
            result['fixture'] = dict(databases=fixture['names'], new_account=new_login)
            admin.execute('USE [' + fixture['names'][0] + ']')
            admin.execute('CREATE SCHEMA other')
            admin.execute("CREATE TABLE other.Projects (id int PRIMARY KEY, label nvarchar(80)); INSERT other.Projects VALUES (7, 'Other schema')")
            admin.execute('CREATE PROCEDURE other.PreviewOnly AS SELECT 1 AS value')
            admin.execute('CREATE PROCEDURE dbo.sp_helpdiagrams AS SELECT 1 AS value')
            admin.execute('GRANT EXECUTE ON dbo.sp_helpdiagrams TO public')
            admin.execute('USE master')
            service = DatabaseTools(Path(temp)/'tools', ROOT/'database_packages', database=object(), guard=nullcontext)
            app = FastAPI(); app.include_router(router); app.include_router(viewer_router)
            def user(request: Request):
                name = request.headers.get('authorization')
                if not name: raise HTTPException(401)
                return dict(username=name, role='user' if name == 'user' else 'admin')
            app.dependency_overrides[get_current_user] = user
            app.dependency_overrides[get_database_tools] = lambda: service
            try:
                with TestClient(app, client=('127.0.0.1', 1234), headers={'authorization':'admin'}) as client:
                    def call(method, path, body=None, status=200):
                        r = client.request(method, path, json=body) if body is not None else client.request(method, path)
                        assert r.status_code == status, (path, r.status_code, r.text)
                        return r.json() if 'application/json' in r.headers.get('content-type','') else r.content
                    source = dict(id='lab-a', name='Lab A', server=fixture['server'], database=fixture['names'][0],
                                  username=new_login, trust_certificate=True)
                    # Verification fails after commit, with a pooled reader session.
                    admin.execute('USE ['+fixture['names'][0]+']; GRANT INSERT ON dbo.Projects TO public; USE master')
                    failed_login = new_login+'_verification'
                    failed_review = call('POST',BASE+'/sources/access/review',dict(source,id='failed-verification',username=failed_login))
                    failure = call('POST',BASE+'/sources/access/create',dict(token=failed_review['token'],windows_auth=True),400)
                    assert 'INSERT' in failure['detail'] and 'check/remove' not in failure['detail'], failure
                    assert not admin.execute('SELECT 1 FROM sys.server_principals WHERE name=?', failed_login).fetchone()
                    admin.execute('USE ['+fixture['names'][0]+']; REVOKE INSERT ON dbo.Projects FROM public; USE master')
                    review = call('POST', BASE+'/sources/access/review', source)
                    assert review['account']==new_login and 'SELECT' in review['sql'] and 'DENY EXECUTE' in review['sql']
                    assert not admin.execute('SELECT name FROM sys.server_principals WHERE name=?',new_login).fetchone()
                    client.headers['authorization']='other-admin'
                    call('POST',BASE+'/sources/access/create',dict(token=review['token'],windows_auth=True),409)
                    client.headers['authorization']='admin'
                    call('POST',BASE+'/sources/access/create',dict(token=review['token'],windows_auth=True))
                    call('POST',BASE+'/sources/access/create',dict(token=review['token'],windows_auth=True),409)
                    created = service.sources.get('lab-a')
                    assert 'password' not in json.dumps(call('GET',BASE+'/viewer-sources'))
                    with service.sources.open(created) as conn:
                        assert conn.execute('SELECT COUNT(*) FROM dbo.Projects').fetchone()[0] == 2
                        for sql in ["INSERT dbo.Projects VALUES(99,'bad')", 'DELETE dbo.Projects', "UPDATE dbo.Projects SET label='bad'", 'CREATE TABLE dbo.Bad(id int)', 'EXEC dbo.sp_helpdiagrams']:
                            try: conn.execute(sql)
                            except Exception: conn.rollback()
                            else: raise AssertionError('New identity can write')
                    inherited = call('POST',BASE+'/sources',dict(source,id='inherited',username=fixture['login'],password=fixture['password']),400)
                    assert 'EXECUTE' in inherited['detail'] and 'sp_helpdiagrams' in inherited['detail'], inherited
                    admin.execute('USE ['+fixture['names'][0]+']; REVOKE EXECUTE ON dbo.sp_helpdiagrams FROM public; USE master')
                    result['checks'].append('New reader denies inherited public EXECUTE; existing reader with EXECUTE rejected; failed verification removes only newly created login and pooled session')
                    collision = call('POST',BASE+'/sources/access/review',dict(source,id='collision',username=fixture['login']))
                    conflict = call('POST',BASE+'/sources/access/create',dict(token=collision['token'],windows_auth=True),400)
                    assert 'already exists' in conflict['detail'] and 'RobotControl_ReadOnly' in conflict['detail']
                    assert admin.execute('SELECT name FROM sys.server_principals WHERE name=?', fixture['login']).fetchone()
                    denied = call('POST',BASE+'/sources/access/review',dict(source,id='denied',username=new_login+'_denied'))
                    error = call('POST',BASE+'/sources/access/create',dict(token=denied['token'],username=fixture['login'],password=fixture['password']),400)
                    assert 'lacks permission' in error['detail'] and 'create login' in error['detail'], error
                    assert fixture['password'] not in json.dumps(error)
                    assert not admin.execute('SELECT name FROM sys.server_principals WHERE name=?',new_login+'_denied').fetchone()
                    signin = call('POST',BASE+'/sources/access/review',dict(source,id='signin',username=new_login+'_signin'))
                    error = call('POST',BASE+'/sources/access/create',dict(token=signin['token'],username=fixture['login'],password='disposable-wrong-password'),400)
                    assert 'rejected the administrator sign-in' in error['detail'], error
                    assert 'disposable-wrong-password' not in json.dumps(error)
                    bad_login = new_login + '_bad'
                    bad = call('POST',BASE+'/sources/access/review',dict(source,id='bad-database',username=bad_login,database=bad_login))
                    missing = call('POST',BASE+'/sources/access/create',dict(token=bad['token'],windows_auth=True),400)
                    assert 'Cannot open database' in missing['detail'], missing
                    assert not admin.execute('SELECT name FROM sys.server_principals WHERE name=?',bad_login).fetchone()
                    result['checks'].append('Existing login never modified; failed database grant transaction removes new login')
                    result['checks'].append('Duplicate name, insufficient SQL authority, rejected sign-in and missing database have distinct actionable HTTP errors; passwords are absent')
                    result['checks'].append('Reviewed create-account HTTP flow; no creation on review; owner/replay enforced; SELECT works and writes/DDL denied')
                    call('POST',BASE+'/sources',dict(source,id='lab-b',name='Lab B',database=fixture['names'][1],username=fixture['login'],password=fixture['password']))
                    a=call('GET','/api/database/tables?source_id=lab-a')['data']['tables']
                    call('GET','/api/database/tables?source_id=lab-b',status=409)
                    call('PUT',BASE+'/viewer-source',{'source_id':'lab-b'})
                    b=call('GET','/api/database/tables?source_id=lab-b')['data']['tables']
                    assert '[other].[Projects]' in a and '[other].[Projects]' not in b
                    from backend.services.report_sources import ReportSources
                    assert ReportSources(service.root).viewer()['id']=='lab-b'
                    call('DELETE',BASE+'/sources/lab-b',status=409)
                    call('PUT',BASE+'/viewer-source',{'source_id':'lab-a'})
                    rows=call('GET','/api/database/tables/%5Bother%5D.%5BProjects%5D?source_id=lab-a')['data']['rows']
                    assert rows[0]['id']==7
                    procedures=call('GET','/api/database/stored-procedures?source_id=lab-a')['data']['procedures']
                    assert any(p['name']=='[other].[PreviewOnly]' and 'SELECT 1' in p['definition'] for p in procedures)
                    call('GET','/api/database/tables')
                    client.headers['authorization']='user'
                    call('GET','/api/database/tables?source_id=lab-b',status=409)
                    call('PUT',BASE+'/viewer-source',{'source_id':'lab-b'},403)
                    call('POST',BASE+'/sources',source,403)
                    client.headers['authorization']='admin'
                    with TestClient(app, client=('10.2.3.4',5),headers={'authorization':'admin','x-forwarded-for':'127.0.0.1'}) as remote:
                        assert remote.post(BASE+'/sources/access/review',json=source).status_code==403
                    result['checks'].append('Two databases and duplicate schema table names browsed; procedure text visible; no native database needed; missing target blocked')
                    # Use a disposable login with write grants in the disposable DB only.
                    admin.execute('USE [' + fixture['names'][1] + ']')
                    admin.execute('GRANT DELETE TO ['+fixture['login']+']')
                    writer=dict(source,id='writer',name='Disposable writer',database=fixture['names'][1],username=fixture['login'],password=fixture['password'],access='operation')
                    call('POST',BASE+'/sources',writer)
                    call('GET','/api/database/tables?source_id=writer',status=409)
                    call('PUT',BASE+'/packages/culture-history/sources',dict(mappings={'primary':'writer'}),400)
                    manifest=dict(contract_version=2,id='remove-project',name='Remove project',version='1.0.0',libraries=[],tools=[dict(id='remove-project',name='Remove project',kind='operation',preview='handler:preview',entrypoint='handler:run',confirmation_field='id',sources=['primary'],inputs=[dict(name='id',label='Project',type='lookup',required=True,lookup=dict(source='primary',query='SELECT id AS value,label FROM dbo.Projects',value_type='integer'))])])
                    handler = "def preview(context, inputs):\n    row=context.connection.execute('SELECT label FROM dbo.Projects WHERE id=?', inputs['id']).fetchone()\n    return dict(summary='Remove project', details=dict(label=row[0] if row else None))\ndef run(context, inputs):\n    context.connection.execute('DELETE dbo.Projects WHERE id=?', inputs['id'])\n    if inputs['id']==2: raise ValueError('Fixture rollback')\n    return dict(message='Removed')\n"
                    output=io.BytesIO()
                    with zipfile.ZipFile(output,'w') as z:
                        z.writestr('manifest.json',json.dumps(manifest));z.writestr('handler.py',handler)
                    assert client.post(BASE+'/packages',files={'file':('operation.zip',output.getvalue())}).status_code==200
                    exported = call('GET', BASE+'/packages/remove-project/export')
                    assert exported == output.getvalue()
                    with zipfile.ZipFile(io.BytesIO(exported)) as archive:
                        assert set(archive.namelist()) == {'manifest.json', 'handler.py'}
                        assert archive.read('handler.py').decode() == handler
                    # Older installations have no retained upload ZIP.
                    directory = service.catalogue.root/service.catalogue.index['remove-project']['directory']
                    (directory/'.package.zip').unlink()
                    (directory/'cache').mkdir(); (directory/'cache'/'private.json').write_text('{}')
                    with zipfile.ZipFile(io.BytesIO(call('GET', BASE+'/packages/remove-project/export'))) as archive:
                        assert set(archive.namelist()) == {'manifest.json', 'handler.py'}
                    client.headers['authorization']='user'
                    call('GET',BASE+'/packages/remove-project/export',status=403)
                    client.headers['authorization']='admin'
                    call('GET',BASE+'/packages/missing/export',status=404)
                    result['checks'].append('Installed package exports exact original bytes; legacy export excludes cache; unavailable and non-admin exports denied')
                    call('PUT',BASE+'/packages/remove-project/sources',dict(mappings={'primary':'lab-a'},operation_source='writer'))
                    options=call('POST',BASE+'/operations/remove-project/choices/id',dict(inputs={}))
                    assert {x['value'] for x in options['options']}=={1,2}
                    call('POST',BASE+'/operations/remove-project/preview',dict(inputs={'id':999}),400)
                    from backend.services.sqlite_safety import SafetyConflict
                    def blocked(): raise SafetyConflict('Robot is busy')
                    service.guard=blocked
                    call('POST',BASE+'/operations/remove-project/preview',dict(inputs={'id':1}),409)
                    service.guard=nullcontext
                    def preview(i): return call('POST',BASE+'/operations/remove-project/preview',dict(inputs={'id':i}))
                    def execute(p,i): return call('POST',BASE+'/operations/execute',dict(token=p['token'],confirmation=str(i)))
                    p=preview(1); assert fixture['names'][1] in p['target']
                    call('POST',BASE+'/sources',writer)
                    assert execute(p,1)['status']=='error'
                    p=preview(2); assert execute(p,2)['status']=='error'
                    p=preview(1)
                    admin.execute('USE ['+fixture['names'][0]+']; DELETE dbo.Projects WHERE id=1')
                    assert execute(p,1)['status']=='error'
                    admin.execute("INSERT dbo.Projects VALUES(1,N'Yeast Ω'); USE ["+fixture['names'][1]+']')
                    p=preview(1); assert execute(p,1)['status']=='succeeded'; assert execute(p,1)['status']=='succeeded'
                    assert admin.execute('SELECT id FROM dbo.Projects').fetchone()[0]==2
                    result['checks'].append('Writer cannot be used by viewers/reports; changed configuration blocks confirmation; rollback and duplicate execution protected')
                    script="import pandas\ndef helper():\n    from openpyxl.styles import Font\n    import missing_library\n    __import__('math')\nraise RuntimeError('MUST NOT EXECUTE')\n"
                    scan=call('POST',BASE+'/authoring/inspect-python',{'source':script})
                    assert scan['available']==['openpyxl','pandas'] and scan['unavailable']==['missing_library'] and scan['undetermined'] and not scan['compatible']
                    call('POST',BASE+'/authoring/inspect-python',{'source':'def ('},400)
                    example="from openpyxl import Workbook\ndef run(context, inputs):\n    b=Workbook(); b.active.append([inputs['label'],42]); b.save(context.output_dir/'example.xlsx'); return 'example.xlsx'\n"
                    assert call('POST',BASE+'/authoring/inspect-python',{'source':example})['compatible']
                    draft=dict(name='Example',package_id='example',sources=[],mappings={},handler=example,original=example,inputs=[dict(name='label',label='Sample name',type='text')])
                    saved=call('POST',BASE+'/drafts',{'draft':draft})
                    editing=call('GET',BASE+'/drafts/'+saved['id']+'/editing-files')
                    with zipfile.ZipFile(io.BytesIO(editing)) as archive:
                        assert set(archive.namelist()) == {'original.py','handler.py','inputs.json','EDITING.md'}
                        assert archive.read('original.py').decode() == example
                        assert json.loads(archive.read('inputs.json'))['inputs'][0]['name'] == 'label'
                    client.headers['authorization']='other-admin'
                    call('GET',BASE+'/drafts/'+saved['id']+'/editing-files',status=404)
                    client.headers['authorization']='admin'
                    result['checks'].append('Editing download retains original, handler and input contract; another author cannot download it')
                    job=call('POST',BASE+'/drafts/'+saved['id']+'/try',dict(inputs={'label':'Demo'}))
                    deadline=time.monotonic()+30
                    while job['status'] in {'pending','running'}:
                        assert time.monotonic()<deadline
                        time.sleep(.05);job=call('GET',BASE+'/reports/'+job['id'])
                    assert job['status']=='ready',job
                    content=call('GET',BASE+'/reports/'+job['id']+'/download')
                    assert list(openpyxl.load_workbook(io.BytesIO(content)).active.values)==[('Demo',42)]
                    (EVIDENCE/'example.xlsx').write_bytes(content)
                    result['checks'].append('Nested/unavailable/dynamic imports detected without running upload; no-database draft generates correct workbook')
                    # Reopen an installed report while preserving sibling tools/assets.
                    package=io.BytesIO()
                    manifest=dict(contract_version=2,id='editable',name='Lab exports',version='2.1.4',libraries=['openpyxl'],tools=[
                        dict(id='editable-report',name='Plate export',kind='report',entrypoint='exporter:run',sources=['primary'],inputs=[
                            dict(name='project',label='Experiment',type='lookup',lookup=dict(source='primary',query='SELECT id AS value,label FROM dbo.Projects',value_type='integer')),
                            dict(name='plate',label='Plate',type='lookup',lookup=dict(source='primary',query='SELECT id AS value,label FROM dbo.Plates WHERE project=?',parameters=['project'],value_type='integer'))]),
                        dict(id='sibling',name='Other export',kind='report',entrypoint='other:run',sources=[],inputs=[])])
                    export_script="from openpyxl import Workbook\ndef run(context, inputs):\n    b=Workbook(); b.active.append([inputs['project'],inputs['plate']]); b.save(context.output_dir/'plate.xlsx'); return 'plate.xlsx'\n"
                    with zipfile.ZipFile(package,'w') as z:
                        z.writestr('manifest.json',json.dumps(manifest));z.writestr('exporter.py',export_script)
                        z.writestr('other.py',"def run(context, inputs):\n    raise ValueError('Other report unchanged')\n");z.writestr('README.md','Keep this supporting file')
                    assert client.post(BASE+'/packages',files={'file':('edit.zip',package.getvalue())}).status_code==200
                    call('PUT',BASE+'/packages/editable/sources',dict(mappings={'primary':'lab-a'}))
                    edited=call('POST',BASE+'/reports/editable-report/edit',{})
                    assert edited['draft']['version']=='2.1.5' and edited['draft']['tool_id']=='editable-report'
                    assert edited['draft']['handler']==export_script and edited['draft']['mappings']=={'primary':'lab-a'}
                    key=edited['id']
                    options=call('POST',BASE+'/drafts/'+key+'/choices/plate',dict(inputs={'project':1}))
                    assert {x['value'] for x in options['options']}=={11,12}
                    def wait_report(job):
                        deadline=time.monotonic()+40
                        while job['status'] in {'pending','running'}:
                            assert time.monotonic()<deadline,job
                            time.sleep(.1);job=call('GET',BASE+'/reports/'+job['id'])
                        return job
                    bad=wait_report(call('POST',BASE+'/drafts/'+key+'/try',dict(inputs={'project':2,'plate':11})))
                    assert bad['status']=='error' and 'no longer available' in bad['error'],bad
                    good=wait_report(call('POST',BASE+'/drafts/'+key+'/try',dict(inputs={'project':1,'plate':11})))
                    assert good['status']=='ready',good
                    assert list(openpyxl.load_workbook(io.BytesIO(call('GET',BASE+'/reports/'+good['id']+'/download'))).active.values)==[(1,11)]
                    stale=call('POST',BASE+'/reports/editable-report/edit',{})
                    reviewed=call('GET',BASE+'/drafts/'+key+'/review')
                    call('POST',BASE+'/drafts/'+key+'/install',dict(revision=edited['revision'],expected_current=reviewed['current_sha256']))
                    exported=call('GET',BASE+'/packages/editable/export')
                    with zipfile.ZipFile(io.BytesIO(exported)) as z:
                        m=json.loads(z.read('manifest.json'))
                        assert m['name']=='Lab exports' and m['version']=='2.1.5' and len(m['tools'])==2
                        assert z.read('README.md')==b'Keep this supporting file'
                        assert z.read('other.py')==b"def run(context, inputs):\n    raise ValueError('Other report unchanged')\n"
                    call('GET',BASE+'/drafts/'+stale['id']+'/review',status=409)
                    call('DELETE',BASE+'/drafts/'+stale['id'])
                    assert call('GET',BASE+'/packages/editable/export')==exported
                    result['checks'].append('Installed report edit retains IDs, siblings, code, assets and mappings; suggests patch version; dependent choices and forged selection checked in child; Excel correct; stale base denied; discard keeps installation')
                    crash=dict(name='Crash fixture',package_id='crash-fixture',sources=[],inputs=[],handler="import os\ndef run(context, inputs):\n    os._exit(17)\n")
                    crashing=call('POST',BASE+'/drafts',dict(draft=crash))
                    failed=wait_report(call('POST',BASE+'/drafts/'+crashing['id']+'/try',dict(inputs={})))
                    assert failed['status']=='error' and 'stopped unexpectedly' in failed['error'],failed
                    import backend.services.database_tools as execution
                    execution.REPORT_TIMEOUT_SECONDS=1
                    try:
                        hanging=call('POST',BASE+'/drafts',dict(draft={**crash,'package_id':'hang-fixture','handler':'import time\ndef run(context, inputs):\n    time.sleep(30)\n'}))
                        expired=wait_report(call('POST',BASE+'/drafts/'+hanging['id']+'/try',dict(inputs={})))
                        assert expired['status']=='error' and 'limit' in expired['error'],expired
                    finally:
                        execution.REPORT_TIMEOUT_SECONDS=300
                    assert wait_report(call('POST',BASE+'/reports/editable-report',dict(inputs={'project':1,'plate':11})))['status']=='ready'
                    result['checks'].append('Report crash and shortened fixture timeout leave HTTP responsive, release workers and allow another successful report; operation preview busy gate and stale/forged choices denied')
                    result['passed']=True
            finally:
                service.close()
    except Exception:
        result['passed']=False
        result['failure']=traceback.format_exc()
        raise
    finally:
        (EVIDENCE/'workspace-http-results.json').write_text(json.dumps(result,indent=2),encoding='utf-8')

if __name__=='__main__': run()
